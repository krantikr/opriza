"""The chase endpoints, driven the way the screen drives them.

`test_chase.py` covers the resolver and the scoring directly. These go through
HTTP, because the drawer's actions - snooze, wake, mark lost - are only useful
if they actually change what the board returns next time it loads.
"""
from datetime import UTC, datetime, timedelta
from decimal import Decimal

from app.models import QuotationPricing, QuotationStatus
from tests.test_chase import reply  # reuse the helper


def test_board_endpoint_serves_rows(client, db, quoted_quotation):
    quoted_quotation.status = QuotationStatus.SENT
    quoted_quotation.sent_at = datetime.now(UTC) - timedelta(days=6)
    item = quoted_quotation.line_items[0]
    db.add(QuotationPricing(quotation_line_item_id=item.id,
                            quote_rate=Decimal("61000.00"),
                            line_total=Decimal("1220000.00"),
                            quantity_mt=Decimal("20")))
    db.commit()
    reply(db, quoted_quotation, body="Rate is high sir.")

    r = client.get("/api/v1/follow-ups/board")
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["counts"]["YOURS"] == 1
    assert data["value_awaiting_you"] > 0
    row = data["rows"][0]
    assert row["turn"] == "YOURS"
    assert row["next_step"] == "Answer on price"
    assert row["last_actor"] == "They replied"
    assert row["last_at"] is not None
    assert row["factors"]


def test_snooze_and_wake(client, db, quoted_quotation):
    quoted_quotation.status = QuotationStatus.SENT
    quoted_quotation.sent_at = datetime.now(UTC) - timedelta(days=6)
    db.commit()
    until = (datetime.now(UTC) + timedelta(days=5)).date().isoformat()

    r = client.post(f"/api/v1/follow-ups/quotation/{quoted_quotation.id}/snooze",
                    json={"until": until})
    assert r.status_code == 200, r.text
    assert client.get("/api/v1/follow-ups/board").json()["counts"]["SNOOZED"] == 1

    r = client.post(f"/api/v1/follow-ups/quotation/{quoted_quotation.id}/wake")
    assert r.status_code == 200, r.text
    assert client.get("/api/v1/follow-ups/board").json()["counts"]["SNOOZED"] == 0


def test_snooze_rejects_today(client, db, quoted_quotation):
    quoted_quotation.status = QuotationStatus.SENT
    quoted_quotation.sent_at = datetime.now(UTC)
    db.commit()
    r = client.post(f"/api/v1/follow-ups/quotation/{quoted_quotation.id}/snooze",
                    json={"until": datetime.now(UTC).date().isoformat()})
    assert r.status_code == 422


def test_mark_lost_records_the_reason(client, db, quoted_quotation):
    quoted_quotation.status = QuotationStatus.SENT
    quoted_quotation.sent_at = datetime.now(UTC) - timedelta(days=6)
    db.commit()

    r = client.post(f"/api/v1/follow-ups/quotation/{quoted_quotation.id}/lost",
                    json={"reason": "PRICE"})
    assert r.status_code == 200, r.text
    db.refresh(quoted_quotation)
    assert quoted_quotation.lost_reason.value == "PRICE"
    # Settled quotations leave the board entirely.
    assert client.get("/api/v1/follow-ups/board").json()["rows"] == []


# ---------------------------------------------------------------------------
# List serializer fallbacks
#
# The Quotations list valued lines from `StockAvailability.suggested_rate`
# only, so a quotation the calculator had priced showed a dash in the list
# while its own detail pane showed a figure. Same for the material name.
# ---------------------------------------------------------------------------


def test_list_values_a_line_from_its_approved_price(client, db, quoted_quotation):
    """`QuotationPricing` wins: it is what the customer was actually quoted."""
    from decimal import Decimal

    from app.models import QuotationPricing

    item = quoted_quotation.line_items[0]
    item.stock_availability.suggested_rate = None
    db.add(
        QuotationPricing(
            quotation_line_item_id=item.id,
            quote_rate=Decimal("61000.00"),
            line_total=Decimal("1220000.00"),
            quantity_mt=Decimal("20"),
        )
    )
    db.commit()

    r = client.get("/api/v1/quotations")
    assert r.status_code == 200, r.text
    row = next(x for x in r.json()["items"] if x["id"] == quoted_quotation.id)
    assert Decimal(row["total_suggested_value"]) == Decimal("1220000.00")


def test_headline_falls_back_to_the_specification(client, db, quoted_quotation):
    """A line that matched nothing still has a grade and a thickness."""
    item = quoted_quotation.line_items[0]
    item.matched_stock_item_id = None
    item.raw_text = None
    db.commit()

    r = client.get("/api/v1/quotations")
    row = next(x for x in r.json()["items"] if x["id"] == quoted_quotation.id)
    assert row["headline_product"] == "10 mm Plate E250"


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------


def test_export_returns_a_real_workbook(client, db, quoted_quotation):
    """A workbook, not a CSV: traders open these in Excel, and a CSV of
    Indian-format rupee figures loses its number typing on the way in."""
    from io import BytesIO

    from openpyxl import load_workbook

    r = client.get("/api/v1/quotations/export")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert ".xlsx" in r.headers["content-disposition"]

    sheet = load_workbook(BytesIO(r.content)).active
    assert sheet["A1"].value == "Quotation"
    numbers = [c.value for c in sheet["A"][1:]]
    assert quoted_quotation.quotation_number in numbers


def test_export_writes_value_as_a_number(client, db, quoted_quotation):
    """The column has to sum in Excel without being cleaned up first."""
    from decimal import Decimal
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models import QuotationPricing

    db.add(
        QuotationPricing(
            quotation_line_item_id=quoted_quotation.line_items[0].id,
            quote_rate=Decimal("61000.00"),
            line_total=Decimal("1220000.00"),
            quantity_mt=Decimal("20"),
        )
    )
    db.commit()

    r = client.get("/api/v1/quotations/export")
    sheet = load_workbook(BytesIO(r.content)).active
    row = next(r for r in sheet.iter_rows(min_row=2, values_only=True)
               if r[0] == quoted_quotation.quotation_number)
    assert row[6] == 1220000.0
    # int or float both sum in Excel; a *string* would not, which is the
    # failure this guards against.
    assert isinstance(row[6], (int, float))


def test_export_path_is_not_read_as_an_id(client):
    """`/export` is registered before `/{quotation_id}`; if that ordering is
    ever lost this 422s trying to parse "export" as an integer."""
    assert client.get("/api/v1/quotations/export").status_code == 200


def test_export_honours_the_search_filter(client, db, quoted_quotation):
    """What you exported should be what you were looking at."""
    from io import BytesIO

    from openpyxl import load_workbook

    r = client.get("/api/v1/quotations/export?search=nothing-matches-this")
    sheet = load_workbook(BytesIO(r.content)).active
    assert sheet.max_row == 1  # headers only


# ---------------------------------------------------------------------------
# Stale pricing snapshots
#
# A pricing row records the state at the moment it was computed. Once rules
# are published that did not exist then, its stored warnings describe a world
# that is gone - and the calculator was repeating "No pricing rule set is
# published" to a desk whose rules were live and configured.
# ---------------------------------------------------------------------------


def test_pricing_priced_before_the_live_rules_is_flagged_stale(client, db, quoted_quotation):

    from app.models import PricingRuleSet, QuotationPricing

    db.add(
        QuotationPricing(
            quotation_line_item_id=quoted_quotation.line_items[0].id,
            rule_set_id=None,  # priced before anything was published
            warnings=["No pricing rule set is published."],
        )
    )
    db.add(PricingRuleSet(version=1, label="Live", is_active=True))
    db.commit()

    r = client.get(f"/api/v1/quotations/{quoted_quotation.id}/pricing")
    assert r.status_code == 200, r.text
    assert r.json()[0]["rules_stale"] is True


def test_pricing_computed_against_the_live_rules_is_not_stale(client, db, quoted_quotation):
    from app.models import PricingRuleSet, QuotationPricing

    rule_set = PricingRuleSet(version=1, label="Live", is_active=True)
    db.add(rule_set)
    db.flush()
    db.add(
        QuotationPricing(
            quotation_line_item_id=quoted_quotation.line_items[0].id,
            rule_set_id=rule_set.id,
        )
    )
    db.commit()

    r = client.get(f"/api/v1/quotations/{quoted_quotation.id}/pricing")
    assert r.json()[0]["rules_stale"] is False
